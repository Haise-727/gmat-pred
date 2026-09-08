"""
Render the paper to a two-column IEEE-style .docx.

This is the reading copy. The submission copy is the IEEEtran LaTeX source next
to it (docs/paper/render_latex.py), and both are rendered from the same document
model in content.py, so they cannot disagree about what the paper says.

Neither renderer holds any prose or any number of its own. To change wording,
edit content.py. To change a number, re-run the experiment that produces the
artifact it comes from — the number is not typed in anywhere.

Two things this renderer has to do by hand that LaTeX does for free:

  * Citations. content.py writes [@key]; IEEE style wants [1]. Keys are numbered
    in order of first appearance and the reference list is generated from
    references.bib at the end.
  * Wide floats. A full-width table or figure in a two-column document needs the
    column count to drop to one and come back, which in Word means a pair of
    continuous section breaks around the float.

Usage:
    python docs/paper/build_paper.py
    python docs/paper/build_paper.py --out /tmp/draft.docx
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from content import CITE_RE, REF_RE, document                      # noqa: E402

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"
FIGDIR = HERE / "figures"
BIB = HERE / "references.bib"

COL_W = Inches(3.4)      # one IEEE column
FULL_W = Inches(7.0)     # both columns

MARKER_RE = re.compile(f"(?:{CITE_RE.pattern})|(?:{REF_RE.pattern})")

ROMAN = ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X",
         "XI", "XII", "XIII", "XIV", "XV"]


# ── bibliography ─────────────────────────────────────────────────────────────

def parse_bib(path: Path) -> dict[str, dict]:
    """
    A deliberately small BibTeX reader: enough for our own file, not general.

    It handles the one nesting level our entries actually use (braced values
    containing braced groups, e.g. {S{\\'a}nchez}) and ignores everything else.
    """
    entries: dict[str, dict] = {}
    if not path.exists():
        return entries
    text = path.read_text()
    for m in re.finditer(r"@(\w+)\s*\{\s*([^,]+),", text):
        key = m.group(2).strip()
        # Walk forward from the entry header to its matching close brace.
        i, depth = m.end() - 1, 1
        while i + 1 < len(text) and depth:
            i += 1
            if text[i] == "{":
                depth += 1
            elif text[i] == "}":
                depth -= 1
        body = text[m.end():i]
        fields = {}
        for fm in re.finditer(r"(\w+)\s*=\s*\{", body):
            j, d = fm.end() - 1, 1
            while j + 1 < len(body) and d:
                j += 1
                if body[j] == "{":
                    d += 1
                elif body[j] == "}":
                    d -= 1
            fields[fm.group(1).lower()] = body[fm.end():j].strip()
        fields["_type"] = m.group(1).lower()
        entries[key] = fields
    return entries


def detex(s: str) -> str:
    """Undo the LaTeX accent escapes our .bib uses, for plain-text output."""
    repl = {
        r"{\"o}": "ö", r"{\"a}": "ä", r"{\"u}": "ü", r"{\"O}": "Ö",
        r"{\'e}": "é", r"{\'a}": "á", r"{\'i}": "í", r"{\'o}": "ó",
        r"{\'{\i}}": "í", r"{\i}": "ı", r"{\~o}": "õ", r"{\~a}": "ã",
        r"{\c{s}}": "ş", r"{\c{c}}": "ç", r"{\v{s}}": "š",
    }
    for a, b in repl.items():
        s = s.replace(a, b)
    return s.replace("{", "").replace("}", "").replace("\\", "")


def format_authors(raw: str) -> str:
    """'Last, First and Last, First' -> 'F. Last, F. Last' (IEEE order)."""
    names = []
    for person in detex(raw).split(" and "):
        person = person.strip()
        if "," in person:
            last, first = [x.strip() for x in person.split(",", 1)]
        else:
            bits = person.split()
            last, first = bits[-1], " ".join(bits[:-1])
        initials = " ".join(f"{b[0]}." for b in first.split() if b)
        names.append(f"{initials} {last}".strip())
    if len(names) > 6:
        return names[0] + " et al."
    return ", ".join(names)


def format_reference(e: dict) -> str:
    """One IEEE-ish reference string. Not a substitute for IEEEtran.bst."""
    authors = format_authors(e.get("author", ""))
    parts = [f"{authors}," if authors else "",
             f'"{detex(e.get("title", ""))},"']
    if e.get("journal"):
        parts.append(f"{detex(e['journal'])},")
    elif e.get("booktitle"):
        parts.append(f"in {detex(e['booktitle'])},")
    if e.get("volume"):
        parts.append(f"vol. {e['volume']},")
    if e.get("number"):
        parts.append(f"no. {e['number']},")
    if e.get("pages"):
        parts.append(f"pp. {e['pages'].replace('--', '–')},")
    if e.get("year"):
        parts.append(f"{e['year']}.")
    if e.get("doi"):
        parts.append(f"doi: {e['doi']}.")
    return " ".join(parts)


# ── docx plumbing ────────────────────────────────────────────────────────────

def set_columns(section, n: int) -> None:
    """Set the column count on a section. python-docx has no API for this."""
    cols = section._sectPr.xpath("./w:cols")[0]
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), "360")          # ~0.25 in gutter


def setup(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(9.5)
    pf = normal.paragraph_format
    pf.space_after = Pt(0)
    pf.first_line_indent = Inches(0.18)
    pf.line_spacing = 1.0
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size in [("Heading 1", 10), ("Heading 2", 9.5), ("Heading 3", 9.5)]:
        st = doc.styles[name]
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.bold = False
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(8)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.first_line_indent = Inches(0)

    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.62)
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)


def rich(p, text: str, cites: dict, refs: dict, size=9.5) -> None:
    """
    Add `text` to paragraph `p`, resolving [@key] and [#label] markers.

    `cites` is the running key -> number map; `refs` maps a float label to its
    printed number. Both are mutated by the caller as the document is walked.
    """
    last = 0
    for m in MARKER_RE.finditer(text):
        if m.start() > last:
            r = p.add_run(text[last:m.start()])
            r.font.size = Pt(size)
        body = m.group(0)
        if body.startswith("[@"):
            nums = []
            for k in body[2:-1].split(","):
                k = k.strip()
                if k not in cites:
                    cites[k] = len(cites) + 1
                nums.append(cites[k])
            r = p.add_run("[" + ", ".join(str(n) for n in sorted(nums)) + "]")
        else:
            label = body[2:-1]
            kind = "Fig." if label.startswith("fig:") else "Table"
            r = p.add_run(f"{kind} {refs.get(label, '?')}")
        r.font.size = Pt(size)
        last = m.end()
    if last < len(text):
        r = p.add_run(text[last:])
        r.font.size = Pt(size)


def number_floats(doc_blocks) -> dict[str, str]:
    """Assign figure and table numbers in document order, before rendering."""
    refs, nfig, ntab = {}, 0, 0
    for kind, payload in doc_blocks:
        if kind == "figure":
            nfig += 1
            refs[payload["label"]] = str(nfig)
        elif kind == "table":
            ntab += 1
            refs[payload["label"]] = ROMAN[ntab] if ntab < len(ROMAN) else str(ntab)
    return refs


def add_float(doc, wide: bool, render) -> None:
    """
    Run `render` inside a one-column island if `wide`, else render in place.

    Word models a column-count change as a section boundary, so a full-width
    float in a two-column document is: continuous break to 1 column, the float,
    continuous break back to 2.
    """
    if not wide:
        render()
        return
    s = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(s, 1)
    s.left_margin = s.right_margin = Inches(0.62)
    render()
    s2 = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(s2, 2)
    s2.left_margin = s2.right_margin = Inches(0.62)


def render_table(doc, t: dict, num: str, cites, refs) -> None:
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Inches(0)
    cap.paragraph_format.space_before = Pt(6)
    r = cap.add_run(f"TABLE {num}\n")
    r.font.size = Pt(8)
    rich(cap, t["caption"], cites, refs, size=8)

    headers = t["headers"]
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.style = "Table Grid"
    tab.autofit = True

    for i, h in enumerate(headers):
        cell = tab.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(h))
        run.bold = True
        run.font.size = Pt(7.5)
        cell.paragraphs[0].paragraph_format.first_line_indent = Inches(0)

    for ri, row in enumerate(t["rows"]):
        cells = tab.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(7.5)
            run.bold = bool(t.get("bold_first_row") and ri == 0)
            cells[i].paragraphs[0].paragraph_format.first_line_indent = Inches(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def render_figure(doc, f: dict, num: str, cites, refs) -> None:
    png = FIGDIR / f"{f['stem']}.png"
    if png.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.add_run().add_picture(str(png),
                                width=FULL_W if f.get("wide") else COL_W)
    cap = doc.add_paragraph()
    cap.paragraph_format.first_line_indent = Inches(0)
    cap.paragraph_format.space_after = Pt(6)
    r = cap.add_run(f"Fig. {num}.  ")
    r.font.size = Pt(8)
    r.bold = True
    rich(cap, f["caption"], cites, refs, size=8)


def build(out_path: Path) -> None:
    blocks = document()
    meta = next(v for k, v in blocks if k == "title")
    refs = number_floats(blocks)
    cites: dict[str, int] = {}

    doc = Document()
    setup(doc)

    # ── title block, full width ──────────────────────────────────────────────
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    r = p.add_run(meta["title"])
    r.font.size = Pt(20)
    for line, size, italic in [(", ".join(meta["authors"]), 11, False),
                               (meta["affiliation"], 10, True)]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        r = p.add_run(line)
        r.font.size = Pt(size)
        r.italic = italic
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Everything from here on is two-column.
    body = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(body, 2)
    body.left_margin = body.right_margin = Inches(0.62)

    n_sec = 0
    n_sub = 0
    for kind, payload in blocks:
        if kind == "title":
            continue

        elif kind == "abstract":
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Inches(0.18)
            r = p.add_run("Abstract—")
            r.bold = True
            r.font.size = Pt(9)
            rich(p, payload, cites, refs, size=9)

        elif kind == "keywords":
            p = doc.add_paragraph()
            r = p.add_run("Index Terms—")
            r.italic = True
            r.font.size = Pt(9)
            r2 = p.add_run(", ".join(payload))
            r2.font.size = Pt(9)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

        elif kind == "h1":
            n_sec += 1
            n_sub = 0
            h = doc.add_heading("", level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            num = ROMAN[n_sec] if n_sec < len(ROMAN) else str(n_sec)
            r = h.add_run(f"{num}.  {payload}")
            r.font.size = Pt(10)
            r.font.name = BODY_FONT

        elif kind == "h2":
            n_sub += 1
            h = doc.add_heading("", level=2)
            r = h.add_run(f"{chr(64 + n_sub)}.  {payload}")
            r.italic = True
            r.font.size = Pt(9.5)
            r.font.name = BODY_FONT

        elif kind == "p":
            p = doc.add_paragraph()
            rich(p, payload, cites, refs)

        elif kind == "bullets":
            for item in payload:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.space_after = Pt(2)
                rich(p, item, cites, refs, size=9)

        elif kind == "table":
            num = refs[payload["label"]]
            add_float(doc, payload.get("wide", False),
                      lambda t=payload, n=num: render_table(doc, t, n, cites, refs))

        elif kind == "figure":
            num = refs[payload["label"]]
            add_float(doc, payload.get("wide", False),
                      lambda f=payload, n=num: render_figure(doc, f, n, cites, refs))

        elif kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.left_indent = Pt(8)
            r = p.add_run(payload)
            r.font.name = MONO_FONT
            r.font.size = Pt(7.5)

    # ── references, numbered in order of first citation ──────────────────────
    bib = parse_bib(BIB)
    h = doc.add_heading("", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("References")
    r.font.size = Pt(10)
    r.font.name = BODY_FONT

    missing = []
    for key, n in sorted(cites.items(), key=lambda kv: kv[1]):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.space_after = Pt(2)
        entry = bib.get(key)
        if entry is None:
            missing.append(key)
            text = f"[{n}] MISSING BIBTEX ENTRY: {key}"
        else:
            text = f"[{n}] {format_reference(entry)}"
        r = p.add_run(text)
        r.font.size = Pt(8)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"  Saved -> {out_path}")
    print(f"  {len(cites)} citations, {len(refs)} numbered floats")
    if missing:
        print(f"  WARNING: no bib entry for {', '.join(missing)}")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(HERE / "OrbitGuard_paper_draft.docx"))
    args = ap.parse_args()
    build(Path(args.out))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
